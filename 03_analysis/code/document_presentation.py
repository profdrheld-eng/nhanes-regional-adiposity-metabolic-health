"""Apply author-confirmed metadata and 11/9-point publication formatting."""
import re
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
AUTHORS='Steffen Held¹, Florian Micke¹, Manuel Matzka¹, Eduard Isenmann²˒³'
AFFILIATIONS=['¹ Department of Management & Sport, IST-University of Applied Sciences, Dusseldorf, NRW, Germany.','² Department of Molecular and Cellular Sports Medicine, Institute for Cardiovascular Research and Sports Medicine, German Sport University Cologne, Cologne, NRW, Germany.','³ Department of Fitness and Health, IST-University of Applied Sciences, Dusseldorf, NRW, Germany.']
def apply_presentation(doc):
    if doc.paragraphs[0].text!='Obesity Science & Practice':
        journal=doc.paragraphs[0].insert_paragraph_before('Obesity Science & Practice')
        for run in journal.runs: run.italic=True
    if not any(AUTHORS==p.text for p in doc.paragraphs):
        author=next((p for p in doc.paragraphs if '[AUTHOR NAMES' in p.text),None)
        if author is None: author=doc.paragraphs[2].insert_paragraph_before(AUTHORS)
        else: author.text=AUTHORS
        from docx.text.paragraph import Paragraph
        anchor=Paragraph(author._p.getnext(),author._parent)
        for text in AFFILIATIONS: anchor.insert_paragraph_before(text)
    for p in doc.paragraphs:
        heading=p.style.name.startswith(('Heading','Title','Subtitle'))
        if p.text=='Obesity Science & Practice': p.alignment=A.CENTER
        elif heading:p.alignment=A.LEFT
        elif p.text:p.alignment=A.JUSTIFY
        small=bool(re.match(r'^(?:Supplementary )?Table\s+[S\d]|^Figure\s+[S\d]|^Note\.|^Locations refer',p.text))
        for r in p.runs:
            if not heading:r.font.size=Pt(9 if small else 11)
            r.font.color.rgb=RGBColor(0,0,0)
    compat=doc.settings.element.find(qn('w:compat'))
    if compat is None:compat=OxmlElement('w:compat');doc.settings.element.append(compat)
    if compat.find(qn('w:doNotExpandShiftReturn')) is None:compat.append(OxmlElement('w:doNotExpandShiftReturn'))
    for table in doc.tables:
        checklist=table.cell(0,0).text=='Item'
        for row in table.rows:
            prop=row._tr.get_or_add_trPr()
            if prop.find(qn('w:cantSplit')) is None:prop.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                if checklist:
                    mar=cell._tc.get_or_add_tcPr().find(qn('w:tcMar'))
                    if mar is not None:
                        for edge in ('top','bottom'):
                            e=mar.find(qn('w:'+edge))
                            if e is not None:e.set(qn('w:w'),'35')
                for p in cell.paragraphs:
                    if p.alignment in (None,A.JUSTIFY):p.alignment=A.LEFT
                    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
