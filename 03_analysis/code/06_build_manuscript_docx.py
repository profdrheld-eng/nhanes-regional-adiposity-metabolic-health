#!/usr/bin/env python3
"""Build the complete publication manuscript as a visually stable DOCX."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT_DIR = ROOT / "05_manuscript"
SOURCE_DIR = MANUSCRIPT_DIR / "source"
FINAL_DIR = MANUSCRIPT_DIR / "final"
RESULTS = ROOT / "04_outputs" / "results"
TABLES = ROOT / "04_outputs" / "tables"
FIGURES = ROOT / "04_outputs" / "figures"
SOURCE = SOURCE_DIR / "manuscript.md"
OUTPUT = FINAL_DIR / "nhanes_leg_to_trunk_metabolic_manuscript.docx"
FINAL_DIR.mkdir(parents=True, exist_ok=True)

BLACK = RGBColor(0, 0, 0)
LIGHT_GRAY = "F4F6F9"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM = 80
CELL_MARGIN_LEFT_RIGHT = 120


def set_run_font(run, size=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = BLACK
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (
        ("top", CELL_MARGIN_TOP_BOTTOM),
        ("bottom", CELL_MARGIN_TOP_BOTTOM),
        ("start", CELL_MARGIN_LEFT_RIGHT),
        ("end", CELL_MARGIN_LEFT_RIGHT),
    ):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "bottom", "insideH"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6" if edge != "insideH" else "2")
        element.set(qn("w:color"), "000000" if edge != "insideH" else "B7B7B7")
    for edge in ("left", "right", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_table_geometry(table, widths_dxa):
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {TABLE_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def create_numbering_definition(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)

    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=9)


def add_inline_markdown(paragraph, text, size=11, bold=False, italic=False):
    token_pattern = re.compile(r"(\*[^*]+\*|`[^`]+`|\[AUTHOR TO[^\]]+\])")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=size, bold=bold, italic=italic)
        token = match.group(0)
        if token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, bold=bold, italic=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, bold=bold, italic=True)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=size, bold=True, italic=italic)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "FFF2CC")
            run._r.get_or_add_rPr().append(shd)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, bold=bold, italic=italic)


def format_p(value):
    value = float(value)
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def prepare_publication_tables():
    table1 = pd.read_csv(TABLES / "table1_weighted_characteristics.csv")
    table1 = table1[["characteristic", "Men", "Women"]].copy()
    table1.columns = ["Characteristic", "Male", "Female"]

    sensitivity = pd.read_csv(RESULTS / "sensitivity_model_coefficients.csv")
    sensitivity_specs = [
        ("single_count_lipid_medication", "log_ltr_z:sexWomen", "Conservative single-count lipid-medication outcome", "Female-to-male ratio of PRs"),
        ("measured_only_outcome", "log_ltr_z:sexWomen", "Measured-only outcome", "Female-to-male ratio of PRs"),
        ("conventional_mets", "log_ltr_z", "Conventional metabolic syndrome", "Male, exposure slope"),
        ("conventional_mets", "log_ltr_z:sexWomen", "Conventional metabolic syndrome", "Female-to-male ratio of PRs"),
        ("bmi_adjusted", "log_ltr_z:sexWomen", "BMI-adjusted model", "Female-to-male ratio of PRs"),
        ("minimal_covariates", "log_ltr_z:sexWomen", "Minimally adjusted model", "Female-to-male ratio of PRs"),
        ("exclude_2017_2018", "log_ltr_z:sexWomen", "Exclude 2017–2018", "Female-to-male ratio of PRs"),
        ("android_gynoid", "scale(log(ag_ratio))", "Android-to-gynoid ratio", "Male, exposure slope"),
        ("android_gynoid", "scale(log(ag_ratio)):sexWomen", "Android-to-gynoid ratio", "Female-to-male ratio of PRs"),
        ("visceral_adipose_tissue", "scale(vat_kg)", "Visceral adipose tissue", "Male, exposure slope"),
        ("visceral_adipose_tissue", "scale(vat_kg):sexWomen", "Visceral adipose tissue", "Female-to-male ratio of PRs"),
        ("leg_and_trunk_components", "scale(leg_fmi)", "Leg and trunk fat-mass indices", "Male, leg fat-mass index"),
        ("leg_and_trunk_components", "scale(trunk_fmi)", "Leg and trunk fat-mass indices", "Male, trunk fat-mass index"),
        ("leg_and_trunk_components", "scale(leg_fmi):sexWomen", "Leg and trunk fat-mass indices", "Female-to-male ratio, leg"),
        ("leg_and_trunk_components", "sexWomen:scale(trunk_fmi)", "Leg and trunk fat-mass indices", "Female-to-male ratio, trunk"),
    ]
    table3_rows = []
    for analysis, term, label, contrast in sensitivity_specs:
        row = sensitivity[
            (sensitivity.analysis == analysis) & (sensitivity.term == term)
        ].iloc[0]
        digits = 3 if analysis == "single_count_lipid_medication" else 2
        table3_rows.append(
            {
                "Analysis": label,
                "Contrast": contrast,
                "PR or ratio": f"{row.exponentiated:.{digits}f}",
                "95% CI": f"{row.ci_low:.{digits}f} to {row.ci_high:.{digits}f}",
                "p": format_p(row.p_value),
            }
        )
    table3 = pd.DataFrame(table3_rows)

    metrics = pd.read_csv(RESULTS / "ml_temporal_metrics.csv")
    panel_labels = {
        "A_base": "Base",
        "B_bmi": "+ BMI",
        "C_bmi_waist": "+ Waist",
        "D_dxa": "+ DXA",
    }
    model_labels = {
        "penalized_logistic": "Penalized logistic",
        "spline_logistic": "Spline logistic",
        "xgboost": "XGBoost",
    }
    table4 = pd.DataFrame(
        {
            "Panel": metrics.panel.map(panel_labels),
            "Model": metrics.model.map(model_labels),
            "ROC AUC (95% CI)": metrics.apply(
                lambda r: f"{r.weighted_auc:.3f} ({r.auc_ci_low:.3f}–{r.auc_ci_high:.3f})",
                axis=1,
            ),
            "Average precision": metrics.weighted_average_precision.map(lambda x: f"{x:.3f}"),
            "Log loss": metrics.weighted_log_loss.map(lambda x: f"{x:.3f}"),
            "Brier score": metrics.weighted_brier.map(lambda x: f"{x:.3f}"),
            "Calibration slope": metrics.calibration_slope.map(lambda x: f"{x:.3f}"),
        }
    )

    components = pd.read_csv(RESULTS / "secondary_component_sex_contrasts.csv")
    outcome_labels = {
        "high_tg": "Elevated triglycerides",
        "low_hdl": "Reduced HDL cholesterol",
        "high_bp": "Elevated blood pressure",
        "high_glucose": "Elevated fasting glucose",
    }
    table_s1 = pd.DataFrame(
        {
            "Outcome": components.outcome.map(outcome_labels),
            "Contrast": components.contrast.replace(
                {"Men": "Male", "Women": "Female", "Interaction": "Female-to-male ratio"}
            ),
            "PR or ratio": components.prevalence_ratio.map(lambda x: f"{x:.2f}"),
            "95% CI": components.apply(
                lambda r: f"{r.ci_low:.2f} to {r.ci_high:.2f}", axis=1
            ),
            "p": components.p_value.map(format_p),
            "FDR-adjusted p": components.p_value_fdr.map(
                lambda x: "—" if pd.isna(x) else format_p(x)
            ),
        }
    )

    incremental = pd.read_csv(RESULTS / "ml_incremental_value.csv")
    table_s2 = pd.DataFrame(
        {
            "Model": incremental.model.map(model_labels),
            "Δ ROC AUC (95% CI)": incremental.apply(
                lambda r: (
                    f"{r.delta_auc_dxa_minus_bmi_waist:.3f} "
                    f"({r.delta_auc_ci_low:.3f}–{r.delta_auc_ci_high:.3f})"
                ),
                axis=1,
            ),
            "Δ Brier score": incremental.delta_brier_dxa_minus_bmi_waist.map(
                lambda x: f"{x:.3f}"
            ),
            "Δ log loss": incremental.delta_log_loss_dxa_minus_bmi_waist.map(
                lambda x: f"{x:.3f}"
            ),
        }
    )

    output_tables = {
        "table1": table1,
        "table3": table3,
        "table4": table4,
        "table_s1": table_s1,
        "table_s2": table_s2,
    }
    for key, frame in output_tables.items():
        frame.to_csv(TABLES / f"{key}_docx.csv", index=False)
    return output_tables


TABLE_META = {
    "table1": {
        "label": "Table 1",
        "title": "Survey-weighted characteristics of the primary analysis population by sex",
        "note": (
            "Note. Continuous variables are weighted arithmetic mean (weighted SD). "
            "Categorical variables are weighted percentages unless stated otherwise. "
            "Counts are unweighted. GED, General Educational Development; SD, standard deviation."
        ),
        "widths": [4680, 2340, 2340],
        "font_size": 8.5,
    },
    "table3": {
        "label": "Table 2",
        "title": "Sensitivity analyses and alternative regional adiposity measures",
        "note": (
            "Note. Main exposure slopes are shown for male participants, the reference category. "
            "Interaction rows are female-to-male ratios of the corresponding PRs. Sample sizes, cases and female slopes are in Supplementary Table S6. Alternative "
            "regional measures were standardized. BMI, body mass index; CI, confidence interval; "
            "PR, prevalence ratio."
        ),
        "widths": [2500, 2600, 1100, 1900, 1260],
        "font_size": 7.5,
    },
    "table4": {
        "label": "Table 3",
        "title": "Performance in the 2017–2018 temporal holdout",
        "note": (
            "Note. Metrics are survey-weighted. ROC AUC intervals are 95% percentile intervals "
            "from 1,000 Rao-Wu rescaled (n_h−1)-PSU replicates within strata, conditional on fitted models. "
            "Base included age, sex, ethnicity, cycle, education, income and smoking. Panels cumulatively add BMI, waist circumference, then total fat-mass index and three regional-fat measures. PSU, primary sampling unit. BMI, body mass index; "
            "CI, confidence interval; DXA, dual-energy X-ray absorptiometry; ROC AUC, area under "
            "the receiver operating characteristic curve."
        ),
        "widths": [700, 1450, 1750, 1160, 1100, 1600, 1600],
        "font_size": 7.5,
    },
    "table_s1": {
        "label": "Supplementary Table S1",
        "title": "Exploratory sex-specific associations with individual metabolic components",
        "note": (
            "Note. Interaction rows are female-to-male ratios of the corresponding PRs. "
            "All component models use n = 3,971 and the primary adjustment set. Component definitions include treatment; generic lipid treatment can make both lipid components positive. Benjamini-Hochberg false-discovery-rate adjustment was applied to the four "
            "interaction tests. CI, confidence interval; FDR, false-discovery rate; "
            "HDL, high-density lipoprotein; PR, prevalence ratio."
        ),
        "widths": [2400, 2100, 1200, 1650, 900, 1110],
        "font_size": 8,
    },
    "table_s2": {
        "label": "Supplementary Table S2",
        "title": "Incremental discrimination from total and regional DXA features beyond BMI and waist circumference",
        "note": (
            "Note. Differences compare the full DXA panel with the BMI-plus-waist panel in the "
            "2017–2018 temporal holdout. The added panel contains total fat-mass index, log leg-to-trunk ratio, android-to-gynoid ratio and visceral adipose tissue mass; its increment cannot be attributed to regional distribution alone. Positive Δ ROC AUC indicates better "
            "discrimination; negative Δ Brier score and Δ log loss indicate improvement. "
            "BMI, body mass index; CI, confidence interval; DXA, dual-energy X-ray absorptiometry; "
            "ROC AUC, area under the receiver operating characteristic curve. Calibration intercept and slope are jointly estimated. Base includes age, sex, ethnicity, cycle, education, income and smoking."
        ),
        "widths": [2400, 3000, 1980, 1980],
        "font_size": 8.5,
    },
}


FIGURE_META = {
    "figure1": {
        "path": FIGURES / "figure1_adjusted_association.png",
        "label": "Figure 1",
        "caption": (
            "Log-linear model: conditional fitted metabolic dysfunction prevalence across the pooled standardized log "
            "leg-to-trunk fat-mass ratio in female and male participants. Lines show model "
            "estimates and colored shaded areas show 95% confidence intervals at the "
            "reference profile (unweighted primary-sample medians for continuous covariates and modes for categories; values in Supplementary Methods S1). Gray bands denote neutral 1-standard-"
            "deviation orientation intervals and are not clinical thresholds. DXA, dual-energy "
            "X-ray absorptiometry; SD, standard deviation."
        ),
        "width": 6.2,
        "alt": "Adjusted metabolic dysfunction prevalence decreases with higher leg-to-trunk fat ratio in female and male participants.",
    },
    "figure2": {
        "path": FIGURES / "figure2_primary_forest.png",
        "label": "Figure 2",
        "caption": (
            "Survey-weighted prevalence ratios and 95% confidence intervals for the primary "
            "sex-specific exposure contrasts and direct interaction contrast. Labels report "
            "the PR or ratio, 95% CI, and two-sided p value. Gray bands provide visual "
            "orientation around the null and are not clinical reference ranges. CI, confidence "
            "interval; PR, prevalence ratio."
        ),
        "width": 6.2,
        "alt": "Forest plot of sex-specific prevalence ratios and the female-to-male interaction ratio.",
    },
    "figure3": {
        "path": FIGURES / "figure3_temporal_model_comparison.png",
        "label": "Figure 3",
        "caption": (
            "Weighted receiver operating characteristic area under the curve in the "
            "2017–2018 temporal test sample for the prespecified feature panels and model "
            "classes. Error bars show 95% intervals from 1,000 bootstrap samples that resampled "
            "primary sampling units within strata with Rao-Wu n_h−1 rescaling, conditional on fitted models. The full DXA panel adds total fat-mass index and three regional-fat measures jointly. Gray bands are 0.10-wide orientation "
            "intervals and do not define performance categories. AUC, area under the curve; BMI, body mass "
            "index; DXA, dual-energy X-ray absorptiometry; ROC, receiver operating characteristic."
        ),
        "width": 6.2,
        "alt": "Temporal-holdout ROC AUC increases after BMI, waist, and regional DXA features are added.",
    },
}


def add_caption(doc, label, title, page_break_before=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run(f"{label}\n")
    set_run_font(run, size=10, bold=True)
    run = paragraph.add_run(title)
    set_run_font(run, size=10, italic=True)


def add_note(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    add_inline_markdown(paragraph, text, size=8.5)


def add_dataframe_table(doc, frame, meta, page_break=True):
    add_caption(
        doc,
        meta["label"],
        meta["title"],
        page_break_before=page_break,
    )
    table = doc.add_table(rows=1, cols=len(frame.columns))
    set_table_geometry(table, meta["widths"])
    set_table_borders(table)
    repeat_header(table.rows[0])
    for idx, column in enumerate(frame.columns):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT if idx in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
        )
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(str(column))
        set_run_font(run, size=meta["font_size"], bold=True)
    for _, values in frame.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if idx in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            text = "" if pd.isna(value) else str(value)
            run = paragraph.add_run(text)
            is_group = idx == 0 and text in {
                "Ethnicity",
                "Education",
                "Smoking status",
            }
            if str(values.iloc[0]) in {"Ethnicity", "Education", "Smoking status"}:
                paragraph.paragraph_format.keep_with_next = True
            set_run_font(run, size=meta["font_size"], bold=is_group)
    set_table_geometry(table, meta["widths"])
    add_note(doc, meta["note"])


def add_figure(doc, meta, page_break=True):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break
    run = paragraph.add_run()
    run.add_picture(str(meta["path"]), width=Inches(meta["width"]))
    drawing = run._r.find(qn("w:drawing"))
    if drawing is not None:
        doc_pr = drawing.find(".//wp:docPr", namespaces=drawing.nsmap)
        if doc_pr is not None:
            doc_pr.set("descr", meta["alt"])
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0
    label_run = caption.add_run(f"{meta['label']}. ")
    set_run_font(label_run, size=9, bold=True)
    add_inline_markdown(caption, meta["caption"], size=9)


def configure_document(doc):
    title_style = doc.styles["Title"]
    title_style.font.color.rgb = BLACK
    for node in list(title_style._element.xpath("./w:pPr/w:pBdr")):
        node.getparent().remove(node)
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = False

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, 18, 10),
        "Heading 2": (13, 12, 6),
        "Heading 3": (12, 8, 4),
        "Heading 4": (11, 6, 3),
    }
    for style_name, (size, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Number", "List Bullet"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.font.color.rgb = BLACK
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    for header_part in (
        section.header,
        section.even_page_header,
        section.first_page_header,
    ):
        header = header_part.paragraphs[0]
        header.text = ""
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = header.add_run("Regional fat distribution and metabolic dysfunction")
        set_run_font(run, size=9)
    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Page ")
    set_run_font(run, size=9)
    add_page_field(footer)

    doc.core_properties.title = (
        "Leg-to-trunk fat distribution and metabolic dysfunction in US adults"
    )
    doc.core_properties.subject = "NHANES 2011–2018 survey-weighted analysis"
    doc.core_properties.author = ""
    doc.core_properties.keywords = (
        "NHANES; DXA; regional adiposity; metabolic dysfunction; sex differences"
    )


def parse_blocks(markdown):
    lines = markdown.splitlines()
    blocks = []
    paragraph_lines = []

    def flush():
        if paragraph_lines:
            blocks.append(("paragraph", " ".join(paragraph_lines).strip()))
            paragraph_lines.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(("heading", level, stripped[level:].strip()))
        elif re.match(r"^\d+\.\s+", stripped):
            flush()
            blocks.append(("number", re.sub(r"^\d+\.\s+", "", stripped)))
        elif stripped.startswith("- "):
            flush()
            blocks.append(("bullet", stripped[2:]))
        else:
            paragraph_lines.append(stripped)
    flush()
    return blocks


def add_body_paragraph(doc, text, style=None):
    paragraph = doc.add_paragraph(style=style)
    if style is None:
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline_markdown(paragraph, text, size=11)
    return paragraph


def build_document():
    tables = prepare_publication_tables()
    markdown = SOURCE.read_text(encoding="utf-8")
    blocks = parse_blocks(markdown)
    title = next(block[2] for block in blocks if block[0] == "heading" and block[1] == 1)

    doc = Document()
    configure_document(doc)

    title_paragraph = doc.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_paragraph.paragraph_format.space_before = Pt(20)
    title_paragraph.paragraph_format.space_after = Pt(12)
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, size=18, bold=True)

    author_paragraph = doc.add_paragraph()
    author_paragraph.paragraph_format.space_after = Pt(4)
    author_run = author_paragraph.add_run("[AUTHOR NAMES AND AFFILIATIONS TO COMPLETE]")
    set_run_font(author_run, size=11, bold=True)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FFF2CC")
    author_run._r.get_or_add_rPr().append(shd)

    running_paragraph = doc.add_paragraph()
    running_paragraph.paragraph_format.space_after = Pt(14)
    running_label = running_paragraph.add_run("Running title: ")
    set_run_font(running_label, size=10, bold=True)
    running_value = running_paragraph.add_run(
        "Regional fat distribution and metabolic dysfunction"
    )
    set_run_font(running_value, size=10)

    skip_running_title = False
    active_numbering_id = None
    in_references = False
    inserted = set()
    for block in blocks:
        if block[0] == "heading":
            active_numbering_id = None
            _, level, text = block
            if level == 1:
                continue
            if text == "Running title":
                skip_running_title = True
                continue
            if skip_running_title:
                skip_running_title = False
                if text != "Abstract":
                    continue
            heading_level = min(level - 1, 4)
            doc.add_heading(text, level=heading_level)
            if text == "References":
                in_references = True
            continue

        text = block[1]
        if skip_running_title:
            skip_running_title = False
            continue
        if block[0] == "number":
            if active_numbering_id is None:
                active_numbering_id = create_numbering_definition(doc)
            paragraph = add_body_paragraph(doc, text)
            apply_numbering(paragraph, active_numbering_id)
        elif block[0] == "bullet":
            active_numbering_id = None
            paragraph = add_body_paragraph(doc, text, style="List Bullet")
        else:
            active_numbering_id = None
            paragraph = add_body_paragraph(doc, text)
        if in_references:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.first_line_indent = Inches(-0.3)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.paragraph_format.space_after = Pt(6)

        if "Weighted age, adiposity, socioeconomic characteristics" in text:
            add_dataframe_table(doc, tables["table1"], TABLE_META["table1"], page_break=False)
            inserted.add("table1")
        if "This contrast describes the imposed log-linear exposure relation" in text:
            add_figure(doc, FIGURE_META["figure1"], page_break=False)
            add_figure(doc, FIGURE_META["figure2"])
            inserted.update({"figure1", "figure2"})
        if "In the joint leg-and-trunk model, the male-reference PRs" in text:
            add_dataframe_table(doc, tables["table3"], TABLE_META["table3"])
            inserted.add("table3")
        if "Figure 3 shows all four panels and three model classes" in text:
            add_figure(doc, FIGURE_META["figure3"])
            inserted.add("figure3")
        if "For penalized logistic regression, the full DXA panel" in text:
            add_dataframe_table(doc, tables["table4"], TABLE_META["table4"])
            inserted.add("table4")

    expected = (set(tables) - {"table_s1", "table_s2"}) | set(FIGURE_META)
    if inserted != expected:
        raise RuntimeError(
            f"Not all manuscript objects were inserted. Missing: {sorted(expected - inserted)}"
        )

    from document_presentation import apply_presentation
    apply_presentation(doc)
    doc.save(OUTPUT)
    print(f"Built {OUTPUT}")
    print(f"Inserted {len(doc.tables)} tables and {len(FIGURE_META)} figures.")


if __name__ == "__main__":
    build_document()
